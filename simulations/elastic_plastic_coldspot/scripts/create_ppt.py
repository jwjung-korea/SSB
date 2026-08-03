import os
from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Title Slide
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "리튬 전착(Electrodeposition) 시뮬레이션 결과"
subtitle.text = "압력 변화(1, 3, 5, 10 MPa)에 따른 응력(증착) 및 소성 변형률(용해) 분석"
title.text_frame.paragraphs[0].font.name = 'Malgun Gothic'
subtitle.text_frame.paragraphs[0].font.name = 'Malgun Gothic'

slide_layout = prs.slide_layouts[5]

ps = [1, 3, 5, 10]
for p in ps:
    slide = prs.slides.add_slide(slide_layout)
    shapes = slide.shapes
    if shapes.title:
        shapes.title.text = f"시뮬레이션 결과: 외부 압력 p = {p} MPa"
        shapes.title.text_frame.paragraphs[0].font.name = 'Malgun Gothic'
    
    img1_path = f"p{p}_t10800_S.png"
    img2_path = f"p{p}_t12960_SDV14.png"
    
    # Left image (Mises Stress)
    if os.path.exists(img1_path):
        left = Inches(0.5)
        top = Inches(1.5)
        width = Inches(5.8)
        shapes.add_picture(img1_path, left, top, width=width)
        
        txBox = shapes.add_textbox(left, Inches(5.8), width, Inches(1.2))
        tf = txBox.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = "Mises 응력 (t=10800s, 최대 증착 시점)"
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.name = 'Malgun Gothic'
        
        p2 = tf.add_paragraph()
        p2.text = f"- 부피 팽창이 최대로 일어났을 때의 응력 분포를 나타냅니다.\n- {p} MPa의 압력 하에서 Cold spot 및 불균일한 표면 주변으로 응력이 집중되는 것을 확인할 수 있습니다."
        p2.font.size = Pt(14)
        p2.font.name = 'Malgun Gothic'
        
    # Right image (Plastic Strain)
    if os.path.exists(img2_path):
        left = Inches(7.0)
        top = Inches(1.5)
        width = Inches(5.8)
        shapes.add_picture(img2_path, left, top, width=width)
        
        txBox2 = shapes.add_textbox(left, Inches(5.8), width, Inches(1.2))
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        p3 = tf2.paragraphs[0]
        p3.text = "등가 소성 변형률 / SDV14 (t=12960s, 최대 용해 시점)"
        p3.font.size = Pt(16)
        p3.font.bold = True
        p3.font.name = 'Malgun Gothic'
        
        p4 = tf2.add_paragraph()
        p4.text = f"- 용해(Stripping) 과정 이후 누적된 비가역적 변형(Dead Li 형성 및 기계적 열화)을 나타냅니다.\n- {p} MPa 압력이 가해질 때 용해 후 소성 변형에 미치는 영향을 보여줍니다."
        p4.font.size = Pt(14)
        p4.font.name = 'Malgun Gothic'

# Last slide for Contact
slide = prs.slides.add_slide(slide_layout)
shapes = slide.shapes
if shapes.title:
    shapes.title.text = "접촉면 및 계면 분석 (t=10800s, 최대 증착 시점)"
    shapes.title.text_frame.paragraphs[0].font.name = 'Malgun Gothic'

contact_img = "interface_comparison_t10800.png"
if os.path.exists(contact_img):
    left = Inches(0.5)
    top = Inches(1.5)
    width = Inches(7.5)
    shapes.add_picture(contact_img, left, top, width=width)
    
    txBox = shapes.add_textbox(Inches(8.2), Inches(2.0), Inches(4.5), Inches(4.0))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = "압력별 계면 특성 비교"
    p0.font.size = Pt(18)
    p0.font.bold = True
    p0.font.name = 'Malgun Gothic'
    
    p1 = tf.add_paragraph()
    p1.text = "1. 접촉 균일성: 외부 압력이 1 MPa에서 10 MPa로 증가함에 따라 계면의 접촉 형태가 더욱 균일해집니다."
    p1.font.size = Pt(16)
    p1.font.name = 'Malgun Gothic'
    
    p2 = tf.add_paragraph()
    p2.text = "2. 국부적 돌기 억제: 높은 압력은 리튬을 소성 변형시켜 국부적인 돌기 성장(수지상 성장 등)을 억제하는 데 도움을 줍니다."
    p2.font.size = Pt(16)
    p2.font.name = 'Malgun Gothic'
    
    p3 = tf.add_paragraph()
    p3.text = "3. 계면 틈새(Gap): 낮은 압력(1, 3 MPa)에서는 10 MPa에 비해 계면에 국부적인 틈새나 심한 불균일성이 발생할 수 있음을 보여줍니다."
    p3.font.size = Pt(16)
    p3.font.name = 'Malgun Gothic'

prs.save('Simulation_Results_Summary_Korean.pptx')
print("Presentation created: Simulation_Results_Summary_Korean.pptx")
